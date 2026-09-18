"""
PRISMA 2020 / PRISMA-trAIce flow-diagram renderer (publication-ready raster + Word).
=====================================================================================
Draws the official PRISMA 2020 study-selection flow — and, when the review is AI-assisted, the **PRISMA-trAIce**
variant (exclusion boxes split into *by Human* / *by AI*, plus the automated-tool-vs-AI footnote) — directly with
Pillow, so there is NO SVG-conversion / headless-browser dependency (only Pillow + python-docx, both already
required). It takes the SAME structured model the web app's `_prisma_model()` produces, so every number is sourced
from `stage_counts.json`; a count that isn't recorded is drawn as a blank "(n = )" exactly like the official
template — never a fabricated 0.

Template followed: PRISMA 2020 new-SR diagram (Page MJ et al. BMJ 2021;372:n71, CC BY 4.0) and the PRISMA-trAIce
extension (https://github.com/cqh4046/PRISMA-trAIce): yellow identification header, blue rotated phase labels,
white count boxes / grey exclusion boxes, side-exit arrows.

Public API:
    render_png(model, project_title="", traice=None, scale=2) -> bytes
    render_jpeg(model, project_title="", traice=None, scale=2, quality=92) -> bytes
    build_docx(model, project_title="", traice=None) -> bytes      # a .docx with the embedded diagram
"""
from __future__ import annotations

import io

from PIL import Image, ImageDraw, ImageFont

# ---- palette (matches the official template) ----
NAVY = (22, 35, 63)
BLACK = (0, 0, 0)
YELLOW = (214, 188, 0)          # the identification header bar
BLUE = (157, 197, 230)          # the rotated phase-label tabs
GREY = (242, 242, 242)          # exclusion / removed boxes
WHITE = (255, 255, 255)
LINE = (60, 60, 60)


def _font(size, bold=False):
    """Best-effort system font; fall back to Pillow's default so this never crashes on a bare machine."""
    for name in ((["arialbd.ttf", "Arialbd.ttf"] if bold else []) + ["arial.ttf", "Arial.ttf",
                 "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"]):
        try:
            return ImageFont.truetype(name, size)
        except Exception:
            continue
    return ImageFont.load_default()


def _n(v):
    """Render a count: a real number (incl. a true 0) as-is, a missing value as a blank — never a fake 0."""
    return str(v) if v is not None and v != "" else ""


def _wrap(draw, text, font, max_w):
    """Greedy word-wrap to a pixel width."""
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=font) <= max_w or not cur:
            cur = t
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def _g(d, *keys):
    for k in keys:
        v = d.get(k)
        if v is not None and v != "":
            return v
    return None


class _Box:
    """A rounded-rectangle box that auto-grows to its content; lines are (text, bold) tuples."""
    def __init__(self, x, y, w, fill=WHITE):
        self.x, self.y, self.w, self.fill = x, y, w, fill
        self.lines = []          # list of (text, bold)

    def add(self, text, bold=False):
        if text is not None:
            self.lines.append((str(text), bold))
        return self

    def draw(self, draw, fonts, pad=12, lh=22):
        # measure (with wrapping) → height
        rendered = []
        inner = self.w - 2 * pad
        for text, bold in self.lines:
            f = fonts["b"] if bold else fonts["r"]
            for ln in _wrap(draw, text, f, inner):
                rendered.append((ln, f))
        h = 2 * pad + max(1, len(rendered)) * lh
        draw.rounded_rectangle([self.x, self.y, self.x + self.w, self.y + h], radius=6,
                               fill=self.fill, outline=LINE, width=2)
        cy = self.y + pad
        for ln, f in rendered:
            draw.text((self.x + self.w / 2, cy), ln, font=f, fill=BLACK, anchor="ma")
            cy += lh
        self.h = h
        return h


def _down_arrow(draw, x, y0, y1):
    draw.line([x, y0, x, y1], fill=LINE, width=2)
    draw.polygon([(x - 6, y1 - 9), (x + 6, y1 - 9), (x, y1)], fill=LINE)


def _right_arrow(draw, x0, x1, y):
    draw.line([x0, y, x1, y], fill=LINE, width=2)
    draw.polygon([(x1 - 9, y - 6), (x1 - 9, y + 6), (x1, y)], fill=LINE)


def _phase_tab(img, draw, label, x, y, h, fonts):
    """A blue vertical phase tab with rotated text (Identification / Screening / Included)."""
    w = 30
    draw.rectangle([x, y, x + w, y + h], fill=BLUE, outline=BLUE)
    txt = Image.new("RGBA", (h, w), (0, 0, 0, 0))
    td = ImageDraw.Draw(txt)
    td.text((h / 2, w / 2), label, font=fonts["b"], fill=NAVY, anchor="mm")
    img.paste(txt.rotate(90, expand=True), (x - 1, y), txt.rotate(90, expand=True))


def _is_traice(model, traice):
    if traice is not None:
        return bool(traice)
    return bool(model.get("traice"))


def _draw(model, project_title="", traice=None, scale=2):
    """Render the whole diagram to a PIL image at `scale`× for crisp raster output."""
    tr = _is_traice(model, traice)
    S = scale
    W = 900 * S
    fonts = {"r": _font(15 * S), "b": _font(15 * S, bold=True),
             "h": _font(19 * S, bold=True), "s": _font(12 * S)}
    # measure pass on a scratch image, then draw on a correctly-sized one
    scratch = Image.new("RGB", (W, 4000 * S), WHITE)
    sd = ImageDraw.Draw(scratch)

    ident = model.get("identification", {})
    rb = ident.get("removed_before_screening", {})
    scr = model.get("screening", {})
    el = model.get("eligibility", {})
    inc = model.get("included", {})

    pad = 14 * S
    tab_w = 30 * S
    col_gap = 26 * S
    left = tab_w + 18 * S
    main_w = int((W - left - 18 * S - col_gap) * 0.5)
    excl_w = W - left - 18 * S - col_gap - main_w
    main_x = left
    excl_x = left + main_w + col_gap

    def newbox(x, w, fill=WHITE):
        return _Box(x, 0, w, fill)

    # --- build the boxes (content only; positions set during layout) ---
    # Identification
    b_id = newbox(main_x, main_w)
    b_id.add("Records identified from:", True)
    src = ident.get("by_source") or {}
    if src:
        for k, v in src.items():
            b_id.add(f"{k} (n = {_n(v)})")
    else:
        b_id.add(f"Databases (n = {_n(_g(ident, 'records_identified'))})")
    b_id.add(f"Registers (n = {_n(ident.get('records_from_registers'))})")

    b_rm = newbox(excl_x, excl_w, GREY)
    b_rm.add("Records removed before screening:", True)
    b_rm.add(f"Duplicate records removed (n = {_n(rb.get('duplicates_removed'))})")
    b_rm.add(f"Records marked as ineligible by automated tools* (n = {_n(rb.get('automation_ineligible'))})")
    b_rm.add(f"Records removed for other reasons (n = {_n(rb.get('removed_other_reasons'))})")

    # Screening
    b_scr = newbox(main_x, main_w)
    b_scr.add("Records screened", True)
    b_scr.add(f"(n = {_n(_g(scr, 'records_screened') if scr.get('records_screened') is not None else ident.get('records_after_dedup'))})")

    b_scr_ex = newbox(excl_x, excl_w, GREY)
    b_scr_ex.add("Records excluded", True)
    if tr:
        b_scr_ex.add(f"by Human (n = {_n(scr.get('excluded_by_human'))})")
        b_scr_ex.add(f"by AI (n = {_n(scr.get('excluded_by_ai'))})")
    else:
        b_scr_ex.add(f"(n = {_n(scr.get('records_excluded'))})")

    b_sought = newbox(main_x, main_w)
    b_sought.add("Reports sought for retrieval", True)
    b_sought.add(f"(n = {_n(scr.get('reports_sought'))})")

    b_notret = newbox(excl_x, excl_w, GREY)
    b_notret.add("Reports not retrieved", True)
    b_notret.add(f"(n = {_n(scr.get('reports_not_retrieved'))})")

    # Eligibility
    b_assess = newbox(main_x, main_w)
    b_assess.add("Reports assessed for eligibility", True)
    b_assess.add(f"(n = {_n(el.get('reports_assessed'))})")

    b_el_ex = newbox(excl_x, excl_w, GREY)
    b_el_ex.add("Reports excluded:", True)
    reasons = el.get("exclusion_reasons") or {}
    rh = el.get("exclusion_reasons_by_human") or {}
    ra = el.get("exclusion_reasons_by_ai") or {}
    if tr and (rh or ra):
        b_el_ex.add("by Human", True)
        for k, v in (rh or {}).items():
            b_el_ex.add(f"   {k} (n = {_n(v)})")
        b_el_ex.add("by AI", True)
        for k, v in (ra or {}).items():
            b_el_ex.add(f"   {k} (n = {_n(v)})")
    elif tr:
        b_el_ex.add(f"by Human (n = {_n(el.get('excluded_by_human'))})")
        b_el_ex.add(f"by AI (n = {_n(el.get('excluded_by_ai'))})")
        for k, v in reasons.items():
            b_el_ex.add(f"   {k} (n = {_n(v)})")
    else:
        for k, v in (reasons or {"(n = " + _n(el.get('reports_excluded')) + ")": None}).items():
            b_el_ex.add(f"{k}" + (f" (n = {_n(v)})" if v is not None else ""))

    # Included
    b_inc = newbox(main_x, main_w)
    b_inc.add("Studies included in review", True)
    b_inc.add(f"(n = {_n(inc.get('studies_included'))})")
    _aw = inc.get("awaiting_classification")
    if _aw not in (None, "", 0, "0"):        # PRISMA 2020: studies awaiting classification (couldn't be obtained/read)
        b_inc.add(f"Studies awaiting classification (n = {_n(_aw)})")
    if inc.get("records_processed_by_ai") is not None:
        b_inc.add(f"Records processed by the AI (n = {_n(inc.get('records_processed_by_ai'))})")

    # --- layout: header, then four rows, then included; phase tabs span their rows ---
    y = pad
    title = (project_title or "").strip()
    # yellow header bar
    hdr_h = 46 * S
    sd.rectangle([left, y, W - 18 * S, y + hdr_h], fill=YELLOW, outline=YELLOW)
    sd.text(((left + W - 18 * S) / 2, y + hdr_h / 2),
            "Identification of studies via databases and registers", font=fonts["h"], fill=BLACK, anchor="mm")
    y += hdr_h + 18 * S

    lh = 24 * S
    rows = []  # (main_box, excl_box_or_None, phase_or_None)
    rows.append((b_id, b_rm, "Identification"))
    rows.append((b_scr, b_scr_ex, "Screening"))
    rows.append((b_sought, b_notret, None))
    rows.append((b_assess, b_el_ex, None))
    rows.append((b_inc, None, "Included"))

    # First pass: compute each row height (max of the two boxes) using a measuring draw.
    positions = []
    yy = y
    for i, (mb, eb, ph) in enumerate(rows):
        mb.y = yy
        h_main = mb.draw(sd, fonts, pad=pad, lh=lh)      # draws on scratch (discarded), sets .h
        h_excl = 0
        if eb is not None:
            eb.y = yy
            h_excl = eb.draw(sd, fonts, pad=pad, lh=lh)
        rh_ = max(h_main, h_excl)
        positions.append((yy, rh_))
        yy += rh_ + 34 * S
    total_h = yy + 70 * S        # room for footnotes

    # Real canvas
    img = Image.new("RGB", (W, int(total_h)), WHITE)
    draw = ImageDraw.Draw(img)
    # header
    yh = pad
    draw.rectangle([left, yh, W - 18 * S, yh + hdr_h], fill=YELLOW, outline=YELLOW)
    draw.text(((left + W - 18 * S) / 2, yh + hdr_h / 2),
              "Identification of studies via databases and registers", font=fonts["h"], fill=BLACK, anchor="mm")

    # phase tabs: Identification spans rows 0; Screening spans rows 1-3; Included spans row 4
    spans = {"Identification": (positions[0][0], positions[0][0] + positions[0][1]),
             "Screening": (positions[1][0], positions[3][0] + positions[3][1]),
             "Included": (positions[4][0], positions[4][0] + positions[4][1])}
    for label, (y0, y1) in spans.items():
        _phase_tab(img, draw, label, 4 * S, int(y0), int(y1 - y0), fonts)

    # boxes + arrows
    for i, (mb, eb, ph) in enumerate(rows):
        y0, rh_ = positions[i]
        mb.y = y0
        mb.draw(draw, fonts, pad=pad, lh=lh)
        cx = main_x + main_w / 2
        if eb is not None:
            eb.y = y0
            eb.draw(draw, fonts, pad=pad, lh=lh)
            _right_arrow(draw, main_x + main_w, excl_x, y0 + rh_ / 2)
        # down arrow to next main box
        if i < len(rows) - 1:
            _down_arrow(draw, cx, y0 + rh_, positions[i + 1][0])

    # footnotes
    fy = positions[-1][0] + positions[-1][1] + 22 * S
    if _is_traice(model, traice):
        for ln in _wrap(draw, "*An automated tool is fundamentally different from an AI application and refers to "
                              "rule-based tools for administrative tasks (e.g. deduplication).", fonts["s"], W - left - 18 * S):
            draw.text((left, fy), ln, font=fonts["s"], fill=(90, 90, 90))
            fy += 18 * S
    draw.text((left, fy), "Source: Page MJ, et al. BMJ 2021;372:n71. doi:10.1136/bmj.n71. Licensed under CC BY 4.0.",
              font=fonts["s"], fill=(90, 90, 90))
    fy += 18 * S
    if title:
        draw.text((left, fy), f"Review: {title}", font=fonts["s"], fill=(90, 90, 90))

    return img


def render_png(model, project_title="", traice=None, scale=2) -> bytes:
    img = _draw(model, project_title, traice, scale)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def render_jpeg(model, project_title="", traice=None, scale=2, quality=92) -> bytes:
    img = _draw(model, project_title, traice, scale).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=quality)
    return buf.getvalue()


def build_docx(model, project_title="", traice=None) -> bytes:
    """A Word document with the rendered PRISMA diagram embedded as a picture (the most reliable way to get the
    publication look into Word without editing the official template's drawing canvas)."""
    import docx
    from docx.shared import Inches
    png = render_png(model, project_title, traice, scale=2)
    doc = docx.Document()
    heading = "PRISMA 2020 flow diagram" + (" (PRISMA-trAIce, AI-assisted)" if _is_traice(model, traice) else "")
    doc.add_heading(heading, level=1)
    if project_title:
        doc.add_paragraph(project_title)
    doc.add_picture(io.BytesIO(png), width=Inches(6.3))
    doc.add_paragraph("Source: Page MJ, et al. BMJ 2021;372:n71. doi:10.1136/bmj.n71. Licensed under CC BY 4.0.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
