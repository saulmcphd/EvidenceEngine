"""
dashboard.py - SUPERSEDED (Phase 3 Streamlit prototype - do not run)
======================================================================
*** THIS FILE IS RETIRED. The real EvidenceEngine app is the webapp/ Flask backend + React frontend, launched
    via "Start EvidenceEngine.bat" -> webapp/serve.py (see CLAUDE.md / README.txt). ***

Why this file still exists but must not be run: it was the original Streamlit prototype for the same 9-step
pipeline, before the webapp/ product reached feature parity and replaced it. It reads and writes the SAME
Outputs/ and okf-bundle/ folders as the live webapp (see EE/BUNDLE below) with none of the webapp's later
safety work (multi-review manager, RAISE provenance gates, reconciliation guardrails, etc.) - running it
against a real review's data risks corrupting or bypassing all of that. It is kept only for reference; do not
extend it, and do not launch it against real data. (Flagged as a genuine confusion risk by the 2026-09 codebase
review - a future session opening this file by name alone could easily mistake it for the current app.)

Original docstring, for reference only (this description of "the dashboard" now describes webapp/, not this file):
A friendly web UI so a non-developer never touches the command line. It WRAPS the tested scripts
(master_records.py, the screeners, reliability.py, okf_tools/okf_writer) - it does not re-implement them.
Pages (sidebar): 1. Setup - project + PICO, criteria.txt, provider/model, API key. 2. Search & Upload - dedupe
-> RIS/CSV/orders downloads. 3. AI screening. 4. Blind screening (instrumented, for the fatigue study).
5. Reliability - recall/F-beta/kappa. 6. Fatigue - mixed model + chart. 7. Report/Export - methods .docx +
BibTeX + bundle .zip. 8. OKF brain - the knowledge graph.
"""

import csv
import io
import json
import os
import re
import subprocess
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import streamlit as st

st.set_page_config(page_title="EvidenceEngine (retired prototype)", page_icon="⚠️")
st.error(
    "**This Streamlit page is retired and must not be used.**\n\n"
    "EvidenceEngine's real app is now the web app - double-click **Start EvidenceEngine.bat** in the "
    "SystematicReview folder (it opens http://localhost:5180 in your browser). This page is an old "
    "prototype that reads/writes the same data folders but has none of the current app's safety checks; "
    "using it on a real review risks corrupting your data.\n\n"
    "If you're a developer who genuinely needs to read this file's old logic for reference, open it in an "
    "editor - do not run it with `streamlit run`."
)
st.stop()

EE = Path(__file__).resolve().parent          # EvidenceEngine/
ROOT = EE.parent                              # repo root
BUNDLE = ROOT / "okf-bundle"
OUT = EE / "Outputs"
OUT.mkdir(exist_ok=True)
sys.path.insert(0, str(EE))

st.set_page_config(page_title="EvidenceEngine", page_icon="🔎", layout="wide")

PROVIDERS = {                                  # display -> a LiteLLM model string the screeners accept
    "Google Gemini (Flash)": "gemini/gemini-2.5-flash",
    "Anthropic Claude (Opus 4.8)": "claude-opus-4-8",
    "OpenAI GPT-4o": "gpt-4o",
    "Local Ollama (Llama 3)": "ollama/llama3",
}
ENV_KEY = {"gemini": "GEMINI_API_KEY", "claude": "ANTHROPIC_API_KEY", "anthropic": "ANTHROPIC_API_KEY",
           "gpt": "OPENAI_API_KEY", "openai": "OPENAI_API_KEY", "ollama": None}


def _env_key_for(model: str) -> str | None:
    for k, v in ENV_KEY.items():
        if model.lower().startswith(k) or f"/{k}" in model.lower() or k in model.lower():
            return v
    return None


# ============================== helpers for the new pages ==============================
# Canonical decision headers — chosen to match screening_import.py's aliases exactly so the
# blind decisions file ingests with no remapping (record_id beats the bare 'id'/'key' aliases).
BLIND_COLS = ["record_id", "human_decision", "decided_at", "screener", "order_index"]


def _now_iso() -> str:
    """Timezone-aware ISO 8601 stamp captured at the MOMENT a blind decision is committed.
    reliability.py parses this for time-on-task; a non-ISO string would silently become NaT."""
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def _append_blind_decision(path: Path, row: dict) -> None:
    """Append ONE blind decision to the shared decisions CSV (crash-safe + resumable):
    each decision is persisted the instant it is made, so a closed browser never loses data."""
    is_new = not path.exists()
    with path.open("a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=BLIND_COLS)
        if is_new:
            w.writeheader()
        w.writerow({k: row.get(k, "") for k in BLIND_COLS})


def _bibtex_from_master(df: "pd.DataFrame") -> str:
    """Hand-rolled BibTeX (no extra dependency) from master_records.csv.
    Field values keep their braces as-is (no paren-mangling of e.g. {DNA}); only the cite-key is sanitised."""
    def esc(v: object) -> str:
        return str(v).strip()
    def citekey(v: object) -> str:
        return re.sub(r"[^A-Za-z0-9_:-]", "", str(v)) or "REC"
    entries = []
    for _, r in df.iterrows():
        rid = esc(r.get("record_id", "")) or "REC"
        authors = esc(r.get("authors", ""))
        for sep in (";", " & ", "|"):
            authors = authors.replace(sep, " and ")
        fields = [("title", esc(r.get("title", ""))), ("author", authors),
                  ("year", esc(r.get("year", ""))), ("doi", esc(r.get("doi", ""))),
                  ("note", f"source: {esc(r.get('source_db', ''))}; EvidenceEngine {rid}")]
        body = ",\n  ".join(f"{k} = {{{v}}}" for k, v in fields if v)
        entries.append(f"@article{{{citekey(rid)},\n  {body}\n}}")
    return "\n\n".join(entries) + "\n"


def _md_to_docx(doc, md_text: str) -> None:
    """Append a markdown string to a python-docx Document (headings / bullets / paragraphs),
    so the canonical raise-disclosure.md flows into the methods .docx without being re-typed."""
    def plain(s: str) -> str:
        return s.replace("**", "").replace("`", "")
    for raw in md_text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(plain(line[4:]).strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(plain(line[3:]).strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(plain(line[2:]).strip(), level=1)
        elif line.lstrip().startswith(("- ", "* ")):
            doc.add_paragraph(plain(line.lstrip()[2:]).strip(), style="List Bullet")
        else:
            doc.add_paragraph(plain(line).strip())


def _load_json(path: Path) -> dict:
    """Read a JSON file, degrading to {} (with a visible warning) if it is missing or malformed,
    so one corrupt file never takes down a whole page."""
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as e:
        st.warning(f"Could not read {path.name} ({e}); treating it as empty. Other items on this page still work.")
        return {}


def _log_consent(screener: str, criteria_text: str) -> None:
    """Append a one-off, auditable consent record (screener · timestamp · criteria fingerprint).
    Recruiting screeners is a human-subjects activity, so consent must leave a trail on disk, not just in session."""
    import hashlib
    path = OUT / "consent_log.csv"
    fp = hashlib.sha1(criteria_text.encode("utf-8")).hexdigest()[:10]
    if path.exists():
        try:
            prior = pd.read_csv(path, dtype=str).fillna("")
            if "screener" in prior.columns and screener in set(prior["screener"]):
                return  # already logged for this screener
        except Exception:
            pass
    is_new = not path.exists()
    with path.open("a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if is_new:
            w.writerow(["screener", "consented_at", "criteria_fingerprint"])
        w.writerow([screener, _now_iso(), fp])


# ============================== sidebar ==============================
st.sidebar.title("🔎 EvidenceEngine")
st.sidebar.caption("AI-assisted systematic review — the AI is a *second* screener; a human reconciles every call.")
page = st.sidebar.radio("Go to", ["Setup", "Search & Upload", "AI screening",
                                  "Blind screening", "Reliability", "Fatigue",
                                  "Report / Export", "OKF brain"])
st.session_state.setdefault("model", PROVIDERS["Google Gemini (Flash)"])
st.session_state.setdefault("blind_screener", None)

# A standing reminder that nothing here is a published result until a real validation run exists.
st.sidebar.divider()
st.sidebar.warning("Demo mode: all metrics are **DEMO/placeholder** until a real blind-human-vs-AI run is loaded.")


# ============================== 1. Setup ==============================
if page == "Setup":
    st.header("1 · Setup & protocol")
    st.write("Tell EvidenceEngine about your review, choose which AI does the screening, and save your API key. "
             "Everything stays on your machine.")

    c1, c2 = st.columns(2)
    with c1:
        st.text_input("Project / review title", key="project")
        st.text_area("Population (P)", key="pico_p", height=70)
        st.text_area("Intervention / Exposure (I/E)", key="pico_i", height=70)
    with c2:
        st.text_area("Comparison (C)", key="pico_c", height=70)
        st.text_area("Outcome (O)", key="pico_o", height=70)

    st.subheader("Eligibility criteria (`criteria.txt`)")
    st.caption("This single file is read by the AI screener. Edit it here and save — no need to open any folder.")
    crit_path = EE / "criteria.txt"
    current = crit_path.read_text(encoding="utf-8") if crit_path.exists() else "REVIEW_TOPIC:\nINCLUSION_CRITERIA:\nEXCLUSION_CRITERIA:\n"
    crit = st.text_area("criteria.txt", value=current, height=260, label_visibility="collapsed")
    if st.button("💾 Save criteria.txt"):
        crit_path.write_text(crit, encoding="utf-8")
        st.success(f"Saved {crit_path}")

    st.subheader("AI provider & API key")
    st.caption("Pick which AI model performs the screening. The key is stored only in a local `.env` file.")
    label = st.selectbox("AI model", list(PROVIDERS), index=0)
    st.session_state["model"] = PROVIDERS[label]
    st.code(st.session_state["model"], language=None)
    need = _env_key_for(st.session_state["model"])
    if need:
        key = st.text_input(f"{need} (your API key)", type="password")
        if st.button("💾 Save API key to .env") and key:
            envp = EE / ".env"
            lines = [l for l in (envp.read_text().splitlines() if envp.exists() else []) if not l.startswith(need + "=")]
            lines.append(f"{need}={key}")
            envp.write_text("\n".join(lines) + "\n", encoding="utf-8")
            st.success(f"Saved {need} to {envp} (kept private; never uploaded).")
    else:
        st.info("Local Ollama needs no API key.")


# ============================== 2. Search & Upload ==============================
elif page == "Search & Upload":
    st.header("2 · Search results → master records")
    st.write("Upload one or more database exports (CSV with columns like *title, abstract, year, authors, doi, "
             "source_db*). EvidenceEngine removes duplicate papers, gives each a stable ID, and prepares the files "
             "you need — then you download them with one click (no folder-digging).")

    n_screeners = st.number_input(
        "Number of blind screeners", min_value=2, max_value=26, value=5,
        help="One separately-randomised order is created per screener for the blind fatigue study. "
             "≥5 enables the proper mixed-effects model (a random intercept per screener); 3–4 falls back "
             "to fixed effects; <3 is descriptive only.")
    seed = st.number_input("Randomisation seed (advanced — keep fixed for reproducibility)",
                           min_value=0, max_value=10_000_000, value=42)

    ups = st.file_uploader("Database export(s) — CSV", type="csv", accept_multiple_files=True)
    if ups and st.button("⚙️ Build master records (dedupe + IDs)"):
        try:
            import master_records as MR
            frames = [pd.read_csv(u).fillna("") for u in ups]
            df = pd.concat(frames, ignore_index=True)
            if hasattr(MR, "normalise_columns"):
                df = MR.normalise_columns(df)
            n_in = len(df)
            deduped, n_dups = (MR.deduplicate(df) if hasattr(MR, "deduplicate") else (df, 0))
            master = MR.build_master(deduped)
            master.to_csv(OUT / "master_records.csv", index=False)
            if hasattr(MR, "write_ris"):
                MR.write_ris(master, OUT / "master_records.ris")
            if hasattr(MR, "write_screening_orders"):
                MR.write_screening_orders(master["record_id"].tolist(), int(n_screeners), int(seed),
                                          OUT / "screening_orders.csv")
            st.session_state["master_built"] = True
            st.success(f"Built {len(master)} master records from {n_in} uploaded rows "
                       f"({n_dups} duplicates removed).")
        except Exception as e:
            st.error(f"Could not build master records: {e}")

    st.subheader("⬇️ Download your files (labelled)")
    st.caption("Each file says what it is and who it's for. The RIS is the one you import into your reference "
               "manager (Mendeley/Zotero/EndNote) for human screening.")
    downloads = [
        ("master_records.ris", "Master records (RIS)", "Import into Mendeley/Zotero/EndNote to screen by hand."),
        ("master_records.csv", "Master records (CSV)", "The spreadsheet the AI screener reads."),
        ("screening_orders.csv", "Per-screener orders (CSV)", "Randomised orders for the blind fatigue study."),
        ("stage_counts.json", "PRISMA counts (JSON)", "Feeds the PRISMA flow diagram."),
    ]
    for fn, label, why in downloads:
        p = OUT / fn
        cols = st.columns([3, 5, 2])
        cols[0].markdown(f"**{label}**")
        cols[1].caption(why)
        if p.exists():
            cols[2].download_button("Download", p.read_bytes(), file_name=fn, key="dl_" + fn)
        else:
            cols[2].caption("— not built yet —")

    with st.expander("Optional: push straight to a reference manager (advanced)"):
        st.caption("If you paste a Zotero or Mendeley API key we can upload the records directly. **If you leave "
                   "these blank, just use the RIS download above** — that's the default and works everywhere.")
        st.text_input("Zotero API key (optional)", type="password")
        st.text_input("Mendeley API key (optional)", type="password")
        st.info("Direct upload is a future option; for now the RIS download is the supported path.")


# ============================== 3. AI screening ==============================
elif page == "AI screening":
    st.header("3 · AI screening (second reviewer)")
    st.write("The AI reads each record and proposes **include / exclude / uncertain** with a reason. It is "
             "recall-first (it over-includes when unsure) and never the final word — you reconcile every call. "
             "Your *human* screening normally happens in your reference manager; upload that RIS in **Reliability** "
             "to compare.")
    stage = st.radio("Stage", ["Title/abstract (5a)", "Full text (5b)"], horizontal=True)
    st.caption(f"Model: `{st.session_state['model']}`  ·  change it on the Setup page.")
    master = OUT / "master_records.csv"
    if not master.exists():
        st.warning("Build your master records first (Search & Upload).")
    else:
        limit = st.number_input("Screen only the first N records (0 = all)", 0, 100000, 0)
        if st.button("▶️ Run AI screening"):
            script = "screener_abstract.py" if stage.startswith("Title") else "screener_fulltext.py"
            cmd = [sys.executable, script, "--model", st.session_state["model"], "--outdir", str(OUT)]
            if stage.startswith("Title"):
                cmd += ["--master", str(master)]
            if limit:
                cmd += ["--limit", str(int(limit))]
            with st.spinner("Screening… (needs the API key saved on Setup)"):
                r = subprocess.run(cmd, cwd=str(EE), capture_output=True, text=True)
            st.code((r.stdout + r.stderr)[-2500:] or "(no output)")
            if r.returncode == 0:
                st.success("Done. See the audit table below.")
            else:
                st.error("The screener reported a problem (often a missing API key).")
        # show the most recent audit CSV
        pat = "Abstract_Audit_*.csv" if stage.startswith("Title") else "FullText_Audit_*.csv"
        files = sorted(OUT.glob(pat))
        if files:
            st.subheader("Latest AI decisions (audit table)")
            st.caption("`Human_Decision` is blank on purpose — fill it from your blind human screen, then reconcile.")
            st.dataframe(pd.read_csv(files[-1]).fillna(""), use_container_width=True, height=360)


# ============================== 4. Reliability ==============================
elif page == "Reliability":
    st.header("5 · Reliability — how good is the AI?")
    st.write("Compare the AI's decisions against your **blind** human decisions (decided *before* seeing the AI). "
             "Recall — the share of relevant studies the AI correctly kept — is the headline, with a 95% confidence "
             "interval. Everything is **DEMO** until you load a real comparison.")

    c1, c2 = st.columns(2)
    human_csv = c1.file_uploader("Blind human decisions (CSV: record_id, human_decision)", type="csv")
    ai_csv = c2.file_uploader("AI audit CSV (record_id, AI_Decision[, AI_Confidence])", type="csv")
    thr = st.slider("Acceptance recall threshold (set a priori, before you look)", 0.50, 1.0, 0.95, 0.01)

    hd = OUT / "human_decisions.csv"
    audits = sorted(OUT.glob("Abstract_Audit_*.csv")) + sorted(OUT.glob("FullText_Audit_*.csv"))
    use_existing = st.checkbox(
        "Use the files I already built (Outputs/human_decisions.csv + latest AI audit)",
        help="Skip re-uploading — use what EvidenceEngine already produced on the Blind screening and AI "
             "screening pages.")
    have_existing = bool(use_existing and hd.exists() and audits)
    if use_existing and not have_existing:
        st.warning("No compiled files found yet — build human_decisions.csv on the Blind screening page and run "
                   "the AI screener first.")

    demo = not have_existing and not (human_csv and ai_csv)
    if demo:
        st.info("No files yet — showing a **DEMO** example (12 relevant, AI misses 1).")
    try:
        import reliability as R
        if demo:
            human = ["include"] * 12 + ["exclude"] * 28
            ai = ["include"] * 11 + ["exclude"] + ["include"] * 3 + ["exclude"] * 25
            m = R.screening_metrics(human, ai, recall_threshold=thr)
        elif have_existing:
            m = R.run_screening(hd, audits[-1], outdir=str(OUT / "reliability"), recall_threshold=thr)
        else:
            hp = OUT / "_human.csv"; ap = OUT / "_ai.csv"
            hp.write_bytes(human_csv.getvalue()); ap.write_bytes(ai_csv.getvalue())
            m = R.run_screening(hp, ap, outdir=str(OUT / "reliability"), recall_threshold=thr)
        tag = " (DEMO)" if demo else ""
        r = m["recall_HEADLINE"]; ci = m["recall_95ci_twosided"]
        a, b, c, d = st.columns(4)
        a.metric(f"Recall{tag}", f"{r:.2f}", help="Share of relevant studies the AI kept. The headline.")
        a.caption(f"95% CI [{ci[0]:.2f}, {ci[1]:.2f}] on {m['positives_in_human']} relevant")
        _beta = m.get("beta", 3)
        b.metric(f"F-beta({_beta:g}){tag}", f"{m['fbeta']:.2f}",
                 help=f"Recall-weighted score (recall ~{_beta ** 2:g}× precision).")
        c.metric(f"Missed (FN){tag}", m["missed_relevant_FN"], help="Relevant studies the AI would have dropped.")
        d.metric(f"Cohen's κ{tag}", f"{m['kappa_secondary']['kappa']:.2f}", help="Agreement beyond chance (secondary).")
        acc = m.get("acceptance", {})
        if acc.get("status") == "estimable":
            ok = acc["passes_headline"]
            (st.success if ok else st.error)(
                f"Acceptance @ recall ≥ {thr:.2f}: **{'PASS' if ok else 'NOT MET'}** "
                f"(judged on the one-sided 95% lower bound = {m['recall_95_onesided_lower']:.2f}, not the point "
                f"estimate). " + (acc.get("precision_warning") or ""))
        elif acc:
            st.warning(f"Acceptance: {acc.get('status')} — {acc.get('reason','')}")
        with st.expander("Methodological caveats (read before quoting any number)"):
            for cav in m.get("_caveats", []):
                st.markdown(f"- {cav}")
    except Exception as e:
        st.error(f"Could not compute metrics: {e}")


# ============================== 5. OKF brain ==============================
elif page == "OKF brain":
    st.header("8 · OKF knowledge brain")
    st.write("EvidenceEngine keeps its whole methodology as an interconnected set of notes (concepts, playbooks, "
             "references). This is the interactive map; you can also download the whole bundle.")
    graph = BUNDLE / "okf-graph.html"
    if graph.exists():
        import streamlit.components.v1 as components
        components.html(graph.read_text(encoding="utf-8"), height=620, scrolling=False)
    else:
        st.warning("Graph not generated yet — run `python okf_tools.py all`.")

    c1, c2 = st.columns(2)
    if c1.button("🧹 Lint the bundle (health check)"):
        r = subprocess.run([sys.executable, "okf_tools.py", "lint"], cwd=str(EE), capture_output=True, text=True)
        c1.code(r.stdout[-1500:])
    if BUNDLE.exists():
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for p in BUNDLE.rglob("*"):
                if p.is_file():
                    z.write(p, p.relative_to(BUNDLE.parent))
        c2.download_button("⬇️ Download the OKF bundle (.zip)", buf.getvalue(),
                           file_name="okf-bundle.zip", mime="application/zip")


# ============================== 6. Blind screening ==============================
elif page == "Blind screening":
    st.header("4 · Blind screening (instrumented)")
    st.write("The **blind** human screening instrument for the validation / fatigue study. You see the eligibility "
             "criteria and one record at a time, and record include / exclude / uncertain **before** any AI output "
             "is shown. Each screener gets their own randomised order, and the moment you decide is timestamped — "
             "that is what lets us measure screening fatigue later, and what keeps the human decision independent "
             "of the AI (no anchoring).")
    st.caption("The AI's calls are hidden here by design. You reconcile against the AI in a later step — never on "
               "this screen.")

    orders_path = OUT / "screening_orders.csv"
    master_path = OUT / "master_records.csv"
    crit_path = EE / "criteria.txt"
    if not (orders_path.exists() and master_path.exists()):
        st.warning("Build your master records and per-screener orders first (Search & Upload).")
    else:
        orders = pd.read_csv(orders_path, dtype=str).fillna("")
        master = pd.read_csv(master_path, dtype=str).fillna("").set_index("record_id")
        screener_cols = [c for c in orders.columns if c.lower() != "position"]
        if not screener_cols:
            st.error("screening_orders.csv has no screener columns — rebuild it on Search & Upload.")
        else:
            who = st.selectbox("Which screener are you?", screener_cols,
                               help="Pick your assigned screener label. Each label has its own randomised order, "
                                    "so list-position is decoupled from database/alphabetical order.")
            st.session_state["blind_screener"] = who

            # Consent / instructions gate — recruiting screeners is a human-subjects activity.
            with st.expander("📋 Instructions & consent (read first)", expanded=True):
                st.markdown(
                    "- Decide **only** from the criteria below and the title/abstract shown.\n"
                    "- Apply the criteria as written; set aside anything you may already know about these papers.\n"
                    "- When genuinely unsure, choose **uncertain** (recall-first — don't exclude on a hunch).\n"
                    "- Your decisions and their timestamps are recorded for a methods study; nothing else is collected.")
                consent = st.checkbox("I have read the instructions and consent to take part.", key=f"consent_{who}")
            if not consent:
                st.info("Tick the consent box above to begin screening.")
            else:
                crit_text = crit_path.read_text(encoding="utf-8") if crit_path.exists() else ""
                _log_consent(who, crit_text)   # auditable on-disk consent record (once per screener)
                with st.expander("📑 Eligibility criteria (criteria.txt) — keep open while you screen", expanded=True):
                    st.code(crit_text or "(criteria.txt not found — set it on the Setup page)", language=None)

                # This screener's worklist, in their own randomised order (position is the order_index).
                worklist = [(int(p), rid) for p, rid in zip(orders["position"], orders[who]) if str(rid).strip()]
                worklist.sort(key=lambda t: t[0])
                # Drop ids that are in the order file but missing from master, so one orphan can't freeze the queue.
                missing = [rid for _, rid in worklist if rid not in master.index]
                worklist = [(p, rid) for p, rid in worklist if rid in master.index]
                total = len(worklist)
                if missing:
                    st.warning(f"{len(missing)} record(s) in your order are missing from master_records.csv and were "
                               f"skipped (rebuild master to restore them): {', '.join(missing[:8])}"
                               + (" …" if len(missing) > 8 else ""))
                if total == 0:
                    st.warning(f"No screenable records are assigned to {who}. Rebuild orders on Search & Upload "
                               "(check the screener count matches the number of screener columns).")
                    st.stop()

                dec_path = OUT / "blind_decisions.csv"
                done = pd.read_csv(dec_path, dtype=str).fillna("") if dec_path.exists() \
                    else pd.DataFrame(columns=BLIND_COLS)
                if len(done):     # protect decided_ids from a stray duplicate (e.g. a second tab)
                    done = done.drop_duplicates(subset=["record_id", "screener"], keep="first")
                mine = done[done["screener"] == who] if len(done) else done
                decided_ids = set(mine["record_id"].astype(str))
                remaining = [(p, rid) for p, rid in worklist if rid not in decided_ids]

                st.progress((total - len(remaining)) / total if total else 0.0,
                            text=f"{total - len(remaining)} of {total} screened ({who})")

                if not remaining:
                    st.success(f"✅ All {total} records screened for {who}. Compile the decisions below, then go to "
                               "Reliability / Fatigue.")
                else:
                    pos, rid = remaining[0]
                    if rid in master.index:
                        rec = master.loc[rid]
                        st.subheader(f"Record {pos} of {total}")
                        st.markdown(f"### {rec.get('title', '') or '(no title)'}")
                        meta = " · ".join([x for x in [rec.get("authors", ""), str(rec.get("year", ""))] if x])
                        if meta:
                            st.caption(meta)
                        st.write(rec.get("abstract", "") or "_(no abstract on file)_")
                        st.caption(f"record_id `{rid}`  ·  your position {pos}")

                        def _commit(decision: str) -> None:
                            if rid in decided_ids:     # idempotent — never double-write a record for a screener
                                st.rerun()
                                return
                            _append_blind_decision(dec_path, {
                                "record_id": rid, "human_decision": decision, "decided_at": _now_iso(),
                                "screener": who, "order_index": pos})
                            st.rerun()

                        b1, b2, b3 = st.columns(3)
                        if b1.button("✅ Include", use_container_width=True):
                            _commit("include")
                        if b2.button("❌ Exclude", use_container_width=True):
                            _commit("exclude")
                        if b3.button("🤔 Uncertain", use_container_width=True):
                            _commit("uncertain")
                    else:
                        st.error(f"{rid} is in the order file but not in master_records.csv — rebuild master records.")

                st.divider()
                with st.expander("Compile decisions for analysis (→ human_decisions.csv)"):
                    st.caption("Combines every screener's blind decisions into the file Reliability and Fatigue "
                               "read. Run it whenever you want to refresh the analysis.")
                    if dec_path.exists() and st.button("🧱 Build human_decisions.csv"):
                        raw = pd.read_csv(dec_path, dtype=str).fillna("")
                        deduped = raw.drop_duplicates(subset=["record_id", "screener"], keep="first")
                        feed = dec_path
                        if len(deduped) < len(raw):
                            feed = OUT / "_blind_dedup.csv"
                            deduped.to_csv(feed, index=False)
                            st.caption(f"Removed {len(raw) - len(deduped)} duplicate (record, screener) row(s) "
                                       "before import, so no screener is double-counted.")
                        cmd = [sys.executable, "screening_import.py", "--input", str(feed), "--kind", "csv",
                               "--master", str(master_path), "--outdir", str(OUT)]
                        r = subprocess.run(cmd, cwd=str(EE), capture_output=True, text=True)
                        st.code((r.stdout + r.stderr)[-1500:] or "(no output)")
                        if r.returncode == 0:
                            st.success(f"Wrote {OUT / 'human_decisions.csv'}.")
                        else:
                            st.error("Import reported a problem (see output above).")


# ============================== 7. Fatigue ==============================
elif page == "Fatigue":
    st.header("6 · Screening fatigue")
    st.write("Does the chance of a screening **error** rise as a screener spends more time on task? The error here "
             "is measured *human-against-human* — each decision is compared to the majority of the **other** "
             "screeners on the same record (leave-one-out), so the AI plays no part in this reference. A positive "
             "time effect whose interval excludes zero is the fatigue signal.")
    try:
        import numpy as np
        import reliability as R
        import plotly.graph_objects as go

        human_path = OUT / "human_decisions.csv"
        real_df, real_note = None, ""
        if human_path.exists():
            try:
                cand = R._fatigue_frame(human_path)
                if len(cand) > 0 and cand["screener"].nunique() >= 2:
                    real_df = cand
                else:
                    real_note = ("Found human_decisions.csv, but it can't support a fatigue model yet "
                                 "(it needs ≥2 blind screeners per record, with per-decision timestamps). ")
            except Exception as ex:
                real_note = f"human_decisions.csv couldn't be read for the fatigue model ({ex}). "
        demo = real_df is None
        if demo:
            if real_note:
                st.warning(real_note)   # real data exists but can't support a model — say so prominently
                st.info("Showing a **DEMO** example below (6 synthetic screeners) — illustration only.")
            else:
                st.info("No instrumented blind decisions yet — showing a **DEMO** example (6 synthetic screeners).")
            rng = np.random.default_rng(7)
            n_s, n_r, rows = 6, 50, []
            for s in range(n_s):
                lab = f"screener_{chr(65 + s)}"
                t = 0.0
                for pos, rec in enumerate(rng.permutation(n_r), start=1):
                    t += float(rng.gamma(2.0, 18.0))                 # variable seconds per record (not a fixed tick)
                    p = 1.0 / (1.0 + np.exp(-(-2.6 + 0.05 * pos)))   # error rate climbs with time on task
                    rows.append({"record_id": f"REC_{rec + 1:04d}", "screener": lab, "order_index": pos,
                                 "cumulative_time": t, "error": int(rng.random() < p)})
            df = pd.DataFrame(rows)
        else:
            df = real_df

        res = R.fatigue_model(df)
        tag = " (DEMO)" if demo else ""

        if res.get("method") == "FAILED":
            st.error("The fatigue model could not be fitted on this data.")
            for k in ("glmm_error", "gee_error", "fit_error"):
                if res.get(k):
                    st.caption(f"{k}: {res[k]}")
        if res.get("warning"):
            st.warning(res["warning"])   # n_screeners < 3 → descriptive only

        a, b, c = st.columns(3)
        a.metric(f"Decisions{tag}", res.get("n_decisions", "—"))
        b.metric(f"Screeners{tag}", res.get("n_screeners", "—"))
        er = res.get("error_rate")
        c.metric(f"Error rate{tag}", f"{er:.2f}" if isinstance(er, (int, float)) else "—",
                 help="Share of decisions that disagreed with the majority of the OTHER screeners on that record.")

        coef = res.get("cumulative_time_coef")
        ci = res.get("cumulative_time_credible_interval") or res.get("cumulative_time_95ci")
        if isinstance(coef, (int, float)) and ci:
            lo, hi = float(ci[0]), float(ci[1])
            detected = bool(res.get("fatigue_detected"))
            colour = "#c0392b" if detected else "#2e7d32"
            fig = go.Figure()
            fig.add_trace(go.Scatter(
                x=[coef], y=["Time-on-task"], mode="markers", marker=dict(size=16, color=colour),
                error_x=dict(type="data", symmetric=False, array=[hi - coef], arrayminus=[coef - lo],
                             thickness=2, width=10, color=colour)))
            fig.add_vline(x=0, line_dash="dash", line_color="#888")
            fig.update_layout(title=f"Effect of time-on-task on screening error{tag}",
                              xaxis_title="log-odds of a screening error per 1 SD of time-on-task "
                                          "(positive = more mistakes the longer you screen)",
                              yaxis_title="", height=240, margin=dict(l=10, r=10, t=50, b=10))
            st.plotly_chart(fig, use_container_width=True)

            verdict = "Fatigue detected" if detected else "No significant fatigue signal"
            if lo > 0:
                zero_msg = "The interval is entirely above 0 → time-on-task credibly increases errors (fatigue)."
            elif hi < 0:
                zero_msg = ("The interval is entirely below 0 → a credible *protective/practice* effect "
                            "(errors fall with time), not fatigue.")
            else:
                zero_msg = "The interval includes 0 → not yet conclusive."
            msg = (f"**{verdict}{tag}.** Time-on-task coefficient {coef:+.3f} "
                   f"(95% interval [{lo:+.3f}, {hi:+.3f}]). {zero_msg}")
            (st.error if detected else st.success)(msg)
            st.caption(f"Model: {res.get('method', '—')} · estimand: {res.get('estimand', '—')} · "
                       f"interval: {res.get('interval_type', '—')}  —  estimand = what the number describes; "
                       "interval = how the uncertainty was computed.")
            op = res.get("order_index_coef")
            if isinstance(op, (int, float)):
                st.caption(f"List-position coefficient (per 1 SD): {op:+.3f}. Time and position are collinear in "
                           "this design, so read them together, not as independent effects.")
        else:
            st.warning("Not enough data to estimate a time-on-task effect yet (need ≥2 screeners per record).")

        d1, d2, d3 = st.columns(3)
        tpc, vif = res.get("time_position_corr"), res.get("time_position_vif")
        d1.metric(f"Time–position r{tag}", f"{tpc:.2f}" if isinstance(tpc, (int, float)) else "—",
                  help="How tangled time-on-task and list-position are (correlation). High = hard to separate.")
        d2.metric(f"VIF{tag}", f"{vif:.1f}" if isinstance(vif, (int, float)) else "—",
                  help="Variance inflation: >10 means time and list-position can't be cleanly separated.")
        d3.metric(f"Distinct orders{tag}", res.get("distinct_orders_across_screeners", "—"),
                  help="How many screeners had a genuinely different order (randomisation sanity check).")
        if res.get("collinearity_warning"):
            st.warning(res["collinearity_warning"])

        with st.expander("Methodological caveats (read before quoting any number)"):
            for cav in res.get("_caveats", []):
                st.markdown(f"- {cav}")
    except Exception as e:
        st.error(f"Could not compute the fatigue model: {e}")


# ============================== 8. Report / Export ==============================
elif page == "Report / Export":
    st.header("7 · Report / Export")
    st.write("Generate the artefacts you paste into your paper: a **Methods .docx** (with the AI-use disclosure "
             "pulled from your canonical RAISE disclosure), a **BibTeX** file of your records, and the OKF bundle. "
             "Every number comes straight from your run — nothing here is invented.")

    master_path = OUT / "master_records.csv"
    counts_path = OUT / "stage_counts.json"
    metrics_path = OUT / "reliability" / "metrics.json"
    disclosure_path = BUNDLE / "raise-disclosure.md"

    counts = _load_json(counts_path)
    metrics = _load_json(metrics_path)

    if metrics:
        st.subheader("Reliability headline")
        rr = metrics.get("recall_HEADLINE")
        ci = metrics.get("recall_95ci_twosided") or [None, None]
        acc = metrics.get("acceptance", {})
        m1, m2, m3 = st.columns(3)
        m1.metric("Recall", f"{rr:.2f}" if isinstance(rr, (int, float)) else "—",
                  help="Headline. Share of relevant studies the AI kept.")
        if isinstance(ci[0], (int, float)):
            m1.caption(f"95% CI [{ci[0]:.2f}, {ci[1]:.2f}]")
        m2.metric("Missed (FN)", metrics.get("missed_relevant_FN", "—"))
        fb = metrics.get("fbeta")
        m3.metric("F-beta", f"{fb:.2f}" if isinstance(fb, (int, float)) else "—")
        if acc.get("status") == "estimable":
            ok = acc.get("passes_headline")
            lb = metrics.get("recall_95_onesided_lower")
            (st.success if ok else st.error)(
                f"Acceptance gate: **{'PASS' if ok else 'NOT MET'}** "
                f"(one-sided 95% lower bound = {lb:.2f})." if isinstance(lb, (int, float)) else
                f"Acceptance gate: **{'PASS' if ok else 'NOT MET'}**.")
    else:
        st.info("No reliability metrics yet — run a comparison on the Reliability page first. The report still "
                "exports your records, PRISMA counts and bundle.")

    if st.button("📝 Generate Methods .docx + BibTeX"):
        try:
            mdf = None
            n_bib = 0
            if master_path.exists():
                mdf = pd.read_csv(master_path, dtype=str).fillna("")
                bib = _bibtex_from_master(mdf)
                (OUT / "references.bib").write_text(bib, encoding="utf-8")
                n_bib = bib.count("@article{")
            # Only assert a step that actually left an artefact on disk — never describe a procedure that didn't run.
            have_human = (OUT / "human_decisions.csv").exists()
            have_ai = bool(sorted(OUT.glob("Abstract_Audit_*.csv")) + sorted(OUT.glob("FullText_Audit_*.csv")))
            have_metrics = isinstance(metrics.get("recall_HEADLINE"), (int, float))
            for key in ("records_after_dedup", "after_dedup", "records_identified", "identified"):
                if key in counts:
                    n_after = counts[key]
                    break
            else:
                n_after = len(mdf) if mdf is not None else "____"
            lines = ["# Methods — AI-assisted study selection (EvidenceEngine)", ""]
            if not (have_human or have_ai):
                lines += ["> DRAFT — no screening artefacts were found in Outputs. The text below is the INTENDED "
                          "design, not a record of what was run; verify every sentence before using it.", ""]
            lines.append(f"After de-duplication, {n_after} records were identified for screening.")
            lines.append("A human reviewer recorded each include/exclude decision **blind** (before seeing any AI "
                         "output), with per-decision timestamps and a per-screener randomised order."
                         if have_human else "Blind human screening: ____ (not yet performed).")
            lines.append("A large language model acted as an independent **second screener** over the same records, "
                         "recall-first (uncertain records were retained)."
                         if have_ai else "AI second-screener run: ____ (not yet performed).")
            if have_human and have_ai:
                lines.append("Disagreements were **reconciled by a human**; the reconciled consensus — not the AI — "
                             "is the review's data.")
            if have_metrics:
                rr = metrics["recall_HEADLINE"]
                ci = metrics.get("recall_95ci_twosided") or [None, None]
                pos = metrics.get("positives_in_human")
                rtxt = f"recall = {rr:.2f}"
                if isinstance(ci[0], (int, float)):
                    rtxt += f", 95% CI [{ci[0]:.2f}, {ci[1]:.2f}]"
                if pos is not None:
                    rtxt += f", on {pos} relevant records"
                acc = metrics.get("acceptance", {})
                if acc.get("sufficient_positives") is False or acc.get("precision_warning") \
                        or metrics.get("rule_of_three_note"):
                    rtxt += (" (small sample — interpret with caution; the CI is dominated by the number of "
                             "relevant records)")
                lines.append(f"Agreement between the AI and the blind human decisions was {rtxt}.")
            else:
                lines.append("Reliability vs the blind human decisions: ____ (run the Reliability comparison).")
            lines += ["", "## AI-use disclosure (RAISE)"]
            preamble = "\n".join(lines) + "\n"
            disclosure_md = disclosure_path.read_text(encoding="utf-8") if disclosure_path.exists() else \
                "_(raise-disclosure.md not generated yet — run okf_writer.write_raise_disclosure.)_"
            full_md = preamble + disclosure_md
            try:
                import docx
                doc = docx.Document()
                _md_to_docx(doc, full_md)
                if counts:
                    doc.add_heading("PRISMA counts", level=2)
                    t = doc.add_table(rows=1, cols=2)
                    t.style = "Table Grid"
                    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Stage", "Count"
                    for k, v in counts.items():
                        if isinstance(v, (int, float, str)):
                            cells = t.add_row().cells
                            cells[0].text, cells[1].text = str(k), str(v)
                doc.save(OUT / "methods.docx")
                made = "methods.docx"
            except ImportError:
                (OUT / "methods.md").write_text(full_md, encoding="utf-8")
                made = "methods.md (python-docx not installed — Markdown fallback)"
            st.success(f"Generated {made} + references.bib ({n_bib} entries) in Outputs."
                       + ("" if n_bib else "  ⚠ No records found — references.bib is empty."))
        except Exception as e:
            st.error(f"Could not generate the report: {e}")

    st.subheader("⬇️ Download")
    downloads = [
        ("methods.docx", "Methods (Word)", "Paste into your paper's Methods; includes the RAISE AI-use disclosure."),
        ("methods.md", "Methods (Markdown fallback)", "Only present if Word export was unavailable."),
        ("references.bib", "References (BibTeX)", "Import into your reference manager / LaTeX."),
        ("master_records.ris", "Records (RIS)", "All records for a reference manager."),
        ("stage_counts.json", "PRISMA counts (JSON)", "Feeds the PRISMA flow diagram."),
        ("human_decisions.csv", "Blind human decisions (CSV)", "The compiled blind screening decisions."),
    ]
    for fn, label, why in downloads:
        p = OUT / fn
        cols = st.columns([3, 5, 2])
        cols[0].markdown(f"**{label}**")
        cols[1].caption(why)
        if p.exists():
            cols[2].download_button("Download", p.read_bytes(), file_name=fn, key="rx_" + fn)
        else:
            cols[2].caption("— not built yet —")

    if BUNDLE.exists():
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for p in BUNDLE.rglob("*"):
                if p.is_file():
                    z.write(p, p.relative_to(BUNDLE.parent))
        st.download_button("⬇️ OKF bundle (.zip)", buf.getvalue(), file_name="okf-bundle.zip",
                           mime="application/zip", key="rx_bundle")
